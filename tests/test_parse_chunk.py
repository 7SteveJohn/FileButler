"""P3 解析/切块测试：PDF 走 PyMuPDF、docx 标题、chunker 标题边界与前缀（不需要 Ollama）。"""
import os
import shutil
import sys
import tempfile

_SANDBOX = tempfile.mkdtemp(prefix="fb_captest_parsechunk_")
os.environ["APPDATA"] = _SANDBOX

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend import db
from backend.knowledge import chunker, parsers

PASS, FAIL = 0, 0


def check(name, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  OK  {name}")
    else:
        FAIL += 1
        print(f"  FAIL {name}  {detail}")


def make_pdf(path):
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "FileButler PDF parsing check.", fontsize=12)
    # 默认 helv 字体不含 CJK（中文会变成占位点），用内置中文字体写入
    page.insert_text((72, 130), "第二行是中文内容，验证双通道解析。",
                     fontsize=12, fontname="china-s")
    doc.save(path)
    doc.close()


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("报销制度", level=1)
    d.add_paragraph("一线城市每晚 500 元。")
    d.add_heading("交通标准", level=2)
    d.add_paragraph("高铁二等座。")
    d.save(path)


def main():
    try:
        db.init_db()

        # --- chunker：标题行是硬边界 + 前缀 ---
        md = "# 报销\n住宿 500 元。\n# 交通\n高铁二等座。"
        chunks = chunker.chunk_text(md, title="制度")
        check("标题行切块数", len(chunks) == 2, str(chunks))
        check("块前缀含文件名与标题", chunks[0].startswith("制度 · 报销"), chunks[0][:30])
        check("前缀正文保留", "住宿 500 元。" in chunks[0], chunks[0])

        # 无标题纯文本：内容不变，仅可有 title 前缀
        plain = "段落一。\n段落二。"
        c2 = chunker.chunk_text(plain, title="笔记")
        check("无标题纯文本单块", len(c2) == 1 and c2[0] == "笔记\n段落一。\n段落二。", str(c2))
        c3 = chunker.chunk_text(plain)
        check("无 title 时无前缀", c3[0] == "段落一。\n段落二。", str(c3))

        # pptx 页标记边界
        ppt = "[第1页]\n第一页内容。\n[第2页]\n第二页内容。"
        c4 = chunker.chunk_text(ppt, title="演示")
        check("页标记切块数", len(c4) == 2, str(c4))

        # --- 解析器 ---
        pdf_path = os.path.join(_SANDBOX, "check.pdf")
        make_pdf(pdf_path)
        text = parsers.parse_file(pdf_path)
        check("PDF 主通道 fitz", bool(text) and "第二行是中文内容" in (text or ""), repr(text)[:80])
        check("PDF 英文行", "FileButler PDF parsing check" in (text or ""))

        docx_path = os.path.join(_SANDBOX, "check.docx")
        make_docx(docx_path)
        dtext = parsers.parse_file(docx_path)
        check("docx 标题行", "# 报销制度" in (dtext or ""), repr(dtext)[:80])
        check("docx 正文", "一线城市每晚 500 元。" in (dtext or ""))

        # --- HTML：剥标签取正文（v3 管线） ---
        html_path = os.path.join(_SANDBOX, "page.htm")
        with open(html_path, "w", encoding="utf-8") as f:
            f.write("<html><head><style>body{color:red}</style>"
                    "<script>var x = 'noise';</script></head><body>"
                    "<h1>报销制度</h1><p>一线城市每晚 500 元。</p>"
                    "<div>高铁二等座。</div></body></html>")
        htext = parsers.parse_file(html_path)
        check("HTML 标题转 md 行", "# 报销制度" in (htext or ""), repr(htext)[:80])
        check("HTML 正文保留", "一线城市每晚 500 元。" in (htext or ""))
        check("HTML script/style 丢弃", "noise" not in (htext or "") and "color" not in (htext or ""),
              repr(htext)[:80])
        check("HTML 标签剥离", "<p>" not in (htext or "") and "<div>" not in (htext or ""))

        # gb18030（中文网页常见编码）也能读
        gbk_path = os.path.join(_SANDBOX, "gbk.htm")
        with open(gbk_path, "wb") as f:
            f.write("<html><body><p>旧版中文页面内容。</p></body></html>".encode("gb18030"))
        gtext = parsers.parse_file(gbk_path)
        check("HTML gb18030 编码", "旧版中文页面内容" in (gtext or ""), repr(gtext)[:80])

        # --- eml：头三行 + text/plain 正文（二进制写入保真 \r\n 行尾） ---
        eml_path = os.path.join(_SANDBOX, "mail.eml")
        with open(eml_path, "wb") as f:
            f.write("From: 张三 <zhangsan@example.com>\r\n"
                    "To: li@example.com\r\n"
                    "Subject: =?utf-8?B?6aG555uu5ZGo5oql?=\r\n"
                    "Date: Fri, 2 Oct 2026 10:00:00 +0800\r\n"
                    "Content-Type: text/plain; charset=utf-8\r\n\r\n"
                    "本周五下班前提交报销单，逾期不候。\r\n".encode("utf-8"))
        etext = parsers.parse_file(eml_path)
        check("eml 主题解码", "项目周报" in (etext or ""), repr(etext)[:100])
        check("eml 发件人", "zhangsan@example.com" in (etext or ""))
        check("eml 正文", "本周五下班前提交报销单" in (etext or ""))
        check("eml 头字段标签", "主题: " in (etext or "") and "发件人: " in (etext or ""))

        # 纯 HTML 邮件正文走剥标签
        html_mail = os.path.join(_SANDBOX, "htmlmail.eml")
        with open(html_mail, "wb") as f:
            f.write("From: a@b.c\r\nSubject: hi\r\n"
                    "Content-Type: text/html; charset=utf-8\r\n\r\n"
                    "<html><body><p>邮件里的正文段落。</p></body></html>\r\n".encode("utf-8"))
        mtext = parsers.parse_file(html_mail)
        check("HTML 邮件剥标签取正文", "邮件里的正文段落" in (mtext or "")
              and "<p>" not in (mtext or ""), repr(mtext)[:100])

        # PARSE_VER 已升到 3（HTML 行为变化触发旧库重解析）
        check("PARSE_VER = 3", parsers.PARSE_VER == "3", parsers.PARSE_VER)
        check("SUPPORTED 含 eml", "eml" in parsers.SUPPORTED)

        print(f"\n{'ALL PASS' if FAIL == 0 else 'FAILURES'} ({PASS} passed, {FAIL} failed)")
    finally:
        db.close_all_conns()
        shutil.rmtree(_SANDBOX, ignore_errors=True)


if __name__ == "__main__":
    main()
    sys.exit(1 if FAIL else 0)
