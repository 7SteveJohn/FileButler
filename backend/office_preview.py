"""Office 文档预览：docx/xlsx/pptx → HTML（供内置预览窗口渲染）。

只提取文字与表格内容，所有文本经 HTML 转义后输出，绝不执行任何脚本；
渲染大文件时限制行/列/页数，保证秒开。
"""
import html


def _esc(s):
    return html.escape(str(s))


def docx_to_html(path):
    """Word：段落（保留标题层级）+ 表格。"""
    from docx import Document
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        # 部分段落（表格内、图片占位等）的 style 可能为 None
        style = ((p.style.name if p.style else "") or "").lower()
        if "heading 1" in style:
            parts.append(f"<h1>{_esc(t)}</h1>")
        elif "heading 2" in style:
            parts.append(f"<h2>{_esc(t)}</h2>")
        elif "heading 3" in style:
            parts.append(f"<h3>{_esc(t)}</h3>")
        elif "title" in style:
            parts.append(f"<h1>{_esc(t)}</h1>")
        else:
            parts.append(f"<p>{_esc(t)}</p>")
    for table in doc.tables[:10]:
        rows = []
        for r in table.rows[:200]:
            cells = [_esc(c.text.strip()) for c in r.cells]
            if any(cells):
                rows.append("<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>")
        if rows:
            parts.append("<table>" + "".join(rows) + "</table>")
    body = "".join(parts) or "<p class='fb-office-empty'>（文档没有可提取的文字内容）</p>"
    return f'<div class="fb-office">{body}</div>'


def xlsx_to_html(path):
    """Excel：每个工作表一张表（最多 200 行 × 30 列，最多 10 个工作表）。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets[:10]:
            rows = []
            for i, row in enumerate(ws.iter_rows(max_row=300, max_col=30)):
                if i >= 200:
                    break
                cells = [_esc(c.value) if c.value is not None else "" for c in row]
                if any(cells):
                    rows.append("<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>")
            if rows:
                parts.append(f"<h3>{_esc(ws.title)}</h3><table>" + "".join(rows) + "</table>")
    finally:
        wb.close()
    body = "".join(parts) or "<p class='fb-office-empty'>（工作簿没有可显示的内容）</p>"
    return f'<div class="fb-office">{body}</div>'


def pptx_to_html(path):
    """PPT：逐页列出形状文本与表格（最多 100 页）。"""
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides):
        if i >= 100:
            break
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(f"<p>{_esc(t)}</p>")
            if getattr(shape, "has_table", False) and shape.has_table:
                for r in shape.table.rows[:50]:
                    cells = [_esc(c.text.strip()) for c in r.cells]
                    if any(cells):
                        texts.append("<p>　" + " ｜ ".join(cells) + "</p>")
        if texts:
            parts.append(f"<h4>第 {i + 1} 页</h4>" + "".join(texts))
    body = "".join(parts) or "<p class='fb-office-empty'>（演示文稿没有可提取的文字内容）</p>"
    return f'<div class="fb-office">{body}</div>'


def to_html(path, ext):
    """按扩展名分发；不支持的返回 None。"""
    ext = ext.lstrip(".").lower()
    if ext == "docx":
        return docx_to_html(path)
    if ext == "xlsx":
        return xlsx_to_html(path)
    if ext == "pptx":
        return pptx_to_html(path)
    return None
